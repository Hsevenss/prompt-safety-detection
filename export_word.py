# -*- coding: utf-8 -*-
"""把项目说明导出成 Word 文档（含流程图与架构图）。"""
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

BASE = r"d:\AITest\aitest\Prompt安全检测数据集"
OUT = f"{BASE}\\项目说明.docx"
FONT = "微软雅黑"


def set_font(run, size=11, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    r = run._element.get_or_add_rPr()
    rfonts = r.rFonts
    rfonts.set(qn("w:eastAsia"), FONT)


def title(doc, text):
    p = doc.add_paragraph()
    p.alignment = 1  # 居中
    set_font(p.add_run(text), size=18, bold=True)


def h(doc, text):
    p = doc.add_paragraph()
    set_font(p.add_run(text), size=14, bold=True)


def body(doc, text):
    p = doc.add_paragraph()
    set_font(p.add_run(text), size=11)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_font(p.add_run(text), size=11)


def make_table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, htext in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        set_font(cell.paragraphs[0].add_run(htext), size=11, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            set_font(cells[i].paragraphs[0].add_run(val), size=11)
    return t


doc = Document()

title(doc, "Prompt 安全检测项目 —— 项目说明")

# 一
h(doc, "一、项目是做什么的")
body(doc, "这是一个提示词安全检测系统，核心任务是判断用户输入的一句话属于「安全(safe)」"
          "还是「不安全(unsafe)」，本质上是一个二分类任务。")

# 二
h(doc, "二、数据规模")
make_table(doc,
           ["数据集", "数量", "说明"],
           [["训练集", "30000 条", "safe 15000 + unsafe 15000，按 1:1 平衡，防止模型偏向某一类"],
            ["测试集", "3195 条", "保留原始分布，用于客观评估"]])

# 三
h(doc, "三、数据预处理（4 步）")
bullet(doc, "只保留中文数据；")
bullet(doc, "去掉长度 ≤5 字的无意义短句；")
bullet(doc, "分层采样，safe/unsafe 各抽 1.5 万条；")
bullet(doc, "统一转成标准训练格式（JSONL）。")

# 四
h(doc, "四、两条技术路线")
body(doc, "微调路线：用 Qwen3-1.7B + LoRA 微调，让模型学会判断安全；微调后在本机启动一个模型接口服务"
          "（http://localhost:8000/v1，仅本机可访问），供评测脚本调用。")
body(doc, "检索路线：用智谱 embedding-3 把训练文本转成向量，建成 FAISS 向量索引，查询时找"
          "「最相似的一条」来判断。")

# 五
h(doc, "五、三种方案的实测数据")
make_table(doc,
           ["方案", "做法", "准确率", "总耗时", "平均耗时"],
           [["方案A 纯微调", "每条都交给模型判断", "85.98%", "7161.62s", "2.24s/条"],
            ["方案B 纯检索", "每条只做向量检索", "75.71%", "739.44s", "0.23s/条"],
            ["方案C 级联", "相似度达标走检索，否则走模型", "85.41%", "6762.68s", "2.12s/条"]])
body(doc, "对比可得：级联比纯微调快了约 5.6%，准确率只低了 0.57 个百分点。")

# 六
h(doc, "六、级联的工作原理")
body(doc, "每条测试样本先检索最近邻，算出一个余弦相似度 s（0~1，越接近 1 表示越像已知样本）：")
body(doc, "s ≥ 阈值 T（如 0.8）：检索很有把握，直接用检索结果（快，不调模型）；")
body(doc, "s < 阈值 T：检索没把握，交给微调模型精判（准，但慢）。")

# 七
h(doc, "七、结论与展望")
body(doc, "实测数据下，纯微调准确率最高（85.98%），但耗时最长；纯检索最快，但准确率偏低。")
body(doc, "级联介于两者之间。从趋势看，随着向量训练数据量增大，检索的命中率和准度会进一步提升，"
          "逐渐追平微调——到那时级联就可以在「速度提升 + 准确率相当」上取得更优的综合效果。")

# 附录：图
h(doc, "附录：流程图与系统架构图")
body(doc, "端到端流程图：")
doc.add_picture(f"{BASE}\\流程图.png", width=Inches(6.0))
doc.add_paragraph()
body(doc, "分层架构图：")
doc.add_picture(f"{BASE}\\系统架构图.png", width=Inches(6.0))

doc.save(OUT)
print("已生成:", OUT)